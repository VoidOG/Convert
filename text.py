import io
from telegram import Update
from fpdf import FPDF
from telegram.ext import CallbackContext
from docx import Document

def handle_text_file(update: Update, context: CallbackContext):
    file = update.message.document
    file.download(file.file_name)
    context.user_data['file_path'] = file.file_name

    # Ask the user what conversion they want to perform
    conversion_options = [
        "Convert to PDF",
        "Convert to DOCX",
        "Convert to TXT"
    ]

    keyboard = [[
        InlineKeyboardButton(option, callback_data=option.replace(" ", "_").lower())
        for option in conversion_options
    ]]
    
    reply_markup = InlineKeyboardMarkup(keyboard)
    update.message.reply_text("Select the conversion you want to perform:", reply_markup=reply_markup)

def convert_text_to_pdf(file_path):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)

    output = io.BytesIO()

    with open(file_path, 'r') as file:
        for line in file:
            pdf.cell(200, 10, txt=line.encode('latin-1', 'replace').decode('latin-1'), ln=True)
    
    pdf.output(output)
    output.seek(0)
    return output

def convert_text_to_docx(file_path):
    doc = Document()
    with open(file_path, 'r') as file:
        for line in file:
            doc.add_paragraph(line)

    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    return output

def convert_text_to_txt(file_path):
    output = io.BytesIO()
    with open(file_path, 'r') as file:
        output.write(file.read().encode('utf-8'))
    output.seek(0)
    return output

def process_conversion(update: Update, context: CallbackContext):
    user_choice = update.callback_query.data.replace("_", " ").title()  # Restore spaces in option
    file_path = context.user_data.get('file_path')

    if user_choice == "Convert To Pdf":
        output = convert_text_to_pdf(file_path)
        filename = "converted.pdf"
    elif user_choice == "Convert To Docx":
        output = convert_text_to_docx(file_path)
        filename = "converted.docx"
    elif user_choice == "Convert To Txt":
        output = convert_text_to_txt(file_path)
        filename = "converted.txt"
    else:
        update.callback_query.answer("Invalid option.")
        return

    update.callback_query.message.reply_document(output, filename=filename)
    update.callback_query.answer()  
